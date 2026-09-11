"""Update BH_CV documentation docx with multi-model pipeline changes."""
import sys
import docx

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

PATH = r"c:\Projects\BH_CV\BH_CV_Compliance_Detector_Documentation.docx"
doc = docx.Document(PATH)
paras = doc.paragraphs


def find(text_fragment):
    for i, p in enumerate(paras):
        if text_fragment in p.text:
            return i
    return -1


def replace_para(idx, new_text):
    if idx < 0:
        return
    p = paras[idx]
    for run in p.runs:
        run.text = ""
    if p.runs:
        p.runs[0].text = new_text
    else:
        p.add_run(new_text)


def clear_para(idx):
    replace_para(idx, "")


# ── Version ───────────────────────────────────────────────────────────────────
i = find("Version 2.0")
replace_para(i, "Version 3.0  |  Multi-Model Handwash Detection  |  Production Release")

# ── Step 1: Adaptive FPS sampling ─────────────────────────────────────────────
i = find("target_fps (def")
replace_para(
    i,
    "FpsSampler uses adaptive dual-rate sampling: no_person_fps=5.0 when no person is "
    "detected (low CPU load), person_fps=10.0 when a person is present (responsive "
    "compliance tracking). Previously a single target_fps was used.",
)

i = find("only the envelope. This keeps queue memory bounded")
replace_para(
    i,
    "The raw frame is NOT sent through the queue — only the lightweight FrameEnvelope. "
    "source_fps (native FPS read from VideoConnector) is now embedded in FrameEnvelope "
    "and propagated through DetectionEvent all the way to the compliance layer for "
    "MM:SS video-time action logging.",
)

# ── Step 2: Three models + source_fps + mp4 ───────────────────────────────────
i = find("StaticModelResolver looks up which models")
replace_para(
    i,
    "StaticModelResolver resolves three models per frame, all nested under the "
    "handwashing use-case group in config.yaml: "
    "handwashing.person (yolo11n.pt, conf=0.25) for person detection and ByteTrack IDs, "
    "handwashing.sink (sink_detection.pt, conf=0.50) for static sink location, "
    "handwashing.hand (handwash_detection.pt, conf=0.50 / presence_conf=0.25) for hand "
    "activity. All three run on every frame.",
)

i = find("YoloDetector runs YOLO inference")
replace_para(
    i,
    "YoloDetector runs YOLO inference for each ModelSpec and converts raw outputs to "
    "typed Detection objects (class_id, class_name, bbox, confidence, use_case). "
    "Detections from all three models are merged into one sorted tuple: "
    "(use_case ASC, class_id ASC, confidence DESC).",
)

i = find("ThumbnailGenerator resizes")
replace_para(
    i,
    "ThumbnailGenerator resizes the frame to max_width=320px and encodes as JPEG "
    "(quality=70). Bounding boxes are drawn only for detections at or above the main "
    "confidence threshold — low-confidence presence detections are excluded from the "
    "visual evidence.",
)

i = find("DetectionEvent to the inference.detections topic. The event is always published")
replace_para(
    i,
    "YoloInferenceEngine publishes a DetectionEvent to inference.detections. The event "
    "carries all merged detections plus source_fps from the FrameEnvelope. Always "
    "published (even when detections is empty) so the compliance state machine can "
    "process absence frames.",
)

# ── Video output: mp4 not avi ─────────────────────────────────────────────────
i = find("the YoloInferenceEngine records an AVI video clip")
replace_para(
    i,
    "When debug.video_saving.enabled is true, YoloInferenceEngine records an MP4 video "
    "clip (codec: mp4v, filename: detected.mp4). The clip is written at the source "
    "video's native FPS so playback speed matches the original footage. Falls back to "
    "fallback_fps=5.0 when native FPS is unknown.",
)

i = find("debug:\n  video_saving:\n    enabled: true\n    fps: 5.0")
replace_para(
    i,
    "debug:\n  video_saving:\n    enabled: true\n"
    "    fallback_fps: 5.0   # used only when source FPS is unknown\n"
    "    filename: detected.mp4",
)

i = find("For confirmed violations, the compliance rule engine also saves a violation evidence clip at evidence/{group_id}/clip.av")
replace_para(
    i,
    "The clip writer is opened lazily on the first detected frame and closed gracefully "
    "on engine.close(). Output: data/result/{session_id}/detected.mp4",
)

# ── Step 3: person-only tracking filter ───────────────────────────────────────
i = find("IoUTracker.update() receives the list of Detection")
replace_para(
    i,
    "ComplianceRuleEngine first filters the merged DetectionEvent to extract only "
    "person detections (class_name == person). Only these are fed to "
    "IoUTracker.update() — sink and hand detections do not need tracking; they are "
    "associated spatially per frame.",
)

# ── Step 4: Observation Building — complete rewrite ───────────────────────────
i = find("ObservationBuilder.build() converts each Track into a domain Observation")
replace_para(
    i,
    "ObservationBuilder.build() receives all detections from all three models plus "
    "active person tracks. It runs the full spatial association pipeline:",
)

i = find("For each track, it checks whether the detection")
replace_para(
    i,
    "1. SPLIT: detections are split by class_name into person / sink / hand lists.",
)

i = find("Water and soap detection are inferred from the full set")
replace_para(
    i,
    "2. HAND->PERSON ASSIGNMENT: For each hand detection, box_center(hand) is tested "
    "against each person bounding box expanded by point_inside_expanded_box "
    "(expand_x=0.15, expand_y=0.20). The nearest qualifying person is assigned that hand.",
)

i = find("inside_sink_zone is set to True for all tracked persons")
replace_para(
    i,
    "3. PERSON->SINK ASSIGNMENT: assign_person_to_sink() finds the nearest sink to the "
    "person lower-center anchor using normalised diagonal distance "
    "(max_person_sink_distance=2.0). No sink assigned if none is within range.",
)

i = find("An Observation is produced for each active track: (camera_id, zone_id, person_id")
replace_para(
    i,
    "4. HAND PROXIMITY: hand_is_near_sink() checks axis-aligned overlap between the "
    "hand box and the sink's expanded proximity zone "
    "(h_margin=0.30, top_margin=1.50, bottom_margin=0.20 x sink dims). "
    "inside_sink_zone=True when any assigned hand is near the sink.\n\n"
    "5. MOVEMENT (EMA): average hand center appended to a rolling deque (size=12). "
    "EMA movement score (alpha=0.2, decay=0.8 when no hands).\n\n"
    "6. WASHING EVIDENCE: has_enough_hands (>=1) AND hand_near_sink AND "
    "(movement_score >= 2.0 when require_hand_movement=true).\n\n"
    "7. TEMPORAL CONFIRMATION: positive_frames increments each frame washing_evidence "
    "is True (reset to 0 otherwise). hands_interacting = positive_frames >= 3.\n\n"
    "8. GRACE PERIOD: hand not near sink but was recently (< 1.0s) -> inside_sink_zone "
    "stays True to absorb brief detection gaps.\n\n"
    "ObservationBuilder is STATEFUL (_tracking dict per person_id). "
    "evict_stale(now, 3.0s) cleans up persons absent too long.",
)

# ── Step 4a: Absence synthesis ────────────────────────────────────────────────
i = find("Because the current model outputs only the")
replace_para(
    i,
    "With multi-model detection, person presence is tracked by the dedicated person "
    "model. When a person is briefly lost (e.g., occlusion), the engine synthesises "
    "absence observations from the StateStore for graceful transitions:",
)

i = find("For each such person, the engine computes frames_lost")
replace_para(
    i,
    "For each person in StateStore NOT seen this frame, "
    "frames_lost = current_frame_id - last_frame_id is computed.",
)

i = find("frames_lost <= MAX_LOST_FRAMES (5): a synthetic Observation is appended with inside_sink_zone=True")
replace_para(
    i,
    "frames_lost <= MAX_LOST_FRAMES (5) AND person not in resting state "
    "(NOT_NEAR_SINK / WASHED / NOT_WASHED): synthetic Observation with "
    "inside_sink_zone=True appended (still at sink).",
)

i = find("If frames_lost > MAX_LOST_FRAMES: a synthetic Observation is appended with inside_sink_zone=False")
replace_para(
    i,
    "frames_lost > MAX_LOST_FRAMES: synthetic Observation with inside_sink_zone=False "
    "appended (person left zone) -> drives state machine toward NOT_NEAR_SINK and "
    "triggers exit-based violation check.",
)

i = find("This synthesis is only applied to persons who were previously known in the StateStore")
replace_para(
    i,
    "Persons already in resting states or with stale state from a previous session "
    "(frames_lost < 0) are skipped.",
)

# ── Step 5: New state names ────────────────────────────────────────────────────
i = find("(or initialised to UNKNOWN if thi")
replace_para(
    i,
    "For each Observation, the current HandwashState is read from the person stored "
    "state (or initialised to NOT_NEAR_SINK if new).",
)

i = find("sequence_valid is False if soap was applied before water was ever detected")
replace_para(
    i,
    "Driving signals from Observation: inside_sink_zone (hand physically at sink, "
    "with grace period) and hands_interacting (temporally confirmed washing). "
    "sequence_valid is always True in the multi-model design.",
)

# ── Step 6: Updated temporal metrics ──────────────────────────────────────────
i = find("at_sink_seconds, water_on_seconds, soap_applied_seconds, washi")
replace_para(
    i,
    "Duration accumulated into at_sink_seconds (NEAR_SINK_NOT_WASHING state) or "
    "washing_seconds (WASHING state). The water_on / soap_applied / rinsing fields "
    "from the previous design have been removed.",
)

i = find("sequence_valid in the person state degrades to False")
replace_para(
    i,
    "observation_positive_frames, observation_no_detection_frames, and "
    "observation_movement_score from the Observation are written into the person "
    "state dict each frame.",
)

i = find("TemporalMetrics is returned: all durations include ongoing time")
replace_para(
    i,
    "TemporalMetrics returned: at_sink_duration, washing_duration (accumulated + "
    "ongoing time in current state), positive_frames, no_detection_frames, "
    "hand_movement_score.",
)

# ── Step 7: Exit-based violation ──────────────────────────────────────────────
i = find("RuleEvaluator.evaluate() receives person_state, TemporalMetrics")
replace_para(
    i,
    "RuleEvaluator.evaluate() receives person_state, TemporalMetrics, camera_id, "
    "zone_id, timestamp, thumbnail, frame_id, and source_fps.",
)

i = find("The evaluator reads current_state from person_state.")
replace_para(
    i,
    "The evaluator reads current_state from person_state and dispatches to one of "
    "three handlers:",
)

i = find("Branch 1")
replace_para(
    i,
    "Branch 1 - Person exited sink zone (NOT_NEAR_SINK): EXIT-BASED COMPLIANCE CHECK. "
    "washing_seconds vs minimum_washing_duration_seconds (5.0s). "
    ">=minimum -> state=WASHED, evidence discarded (compliant, logged as WASHED + MM:SS). "
    "<minimum -> state=NOT_WASHED, evidence FROZEN -> violation (logged as NOT WASHED + MM:SS).",
)

i = find("Branch 2")
replace_para(
    i,
    "Branch 2 - Person confirmed washing (WASHING): any unconfirmed absence evidence "
    "group is discarded. No violation triggered.",
)

i = find("Branch 3")
replace_para(
    i,
    "Branch 3 - Person at sink not confirmed washing (NEAR_SINK_NOT_WASHING): "
    "Evidence group opened after 2 consecutive frames (debounce). Each frame thumbnail "
    "buffered in EvidenceBuffer (ring buffer, max=10). NO violation here - only on "
    "zone exit.",
)

i = find("If no active group: mint a new UUID group_id, record absence_started_at")
replace_para(i, "violation_confirmed=True once set; outcome=None for all subsequent frames in the same episode (one alert per group_id).")

i = find("Add an EvidenceItem (frame_id, timestamp, thumbnail")
replace_para(i, "KEY CHANGE FROM V2: Violations are triggered on ZONE EXIT, not on absence duration exceeding a threshold. The evaluator only fires when the state machine transitions to NOT_NEAR_SINK.")

i = find("Compute absence_duration = (timestamp - absence_started_at)")
replace_para(i, "")

i = find("If absence_duration >= absence_threshold_seconds")
replace_para(i, "")

i = find("Return RuleEvaluationResult(person_state, outcome='violation'")
replace_para(i, "")

i = find("If violation_confirmed is already True (idempotency guard)")
replace_para(i, "")

# ── Section 4.2: State diagram ────────────────────────────────────────────────
i = find("UNKNOWN\n       |  person enters sink zone")
replace_para(
    i,
    "NOT_NEAR_SINK  (initial / resting)\n"
    "    |\n"
    "    | hand reaches sink (inside_sink_zone=True)\n"
    "    v\n"
    "NEAR_SINK_NOT_WASHING\n"
    "    |\n"
    "    | positive_frames >= confirmation_frames (3)\n"
    "    v\n"
    "WASHING\n"
    "    |  washing stops (positive_frames reset) -> NEAR_SINK_NOT_WASHING\n"
    "    |\n"
    "    | zone exit (inside_sink_zone=False)\n"
    "    v\n"
    "NOT_NEAR_SINK  ->  RuleEvaluator checks washing_seconds:\n"
    "    >= 5.0s  ->  WASHED     (compliant)\n"
    "    <  5.0s  ->  NOT_WASHED (violation)\n\n"
    "WASHED / NOT_WASHED are terminal; reset to NOT_NEAR_SINK on next detection.",
)

# ── Section 6.3: ObservationBuilder ──────────────────────────────────────────
i = find("Converts IoUTracker output (list[Track]) plus DetectionEvent metadata")
replace_para(
    i,
    "Stateful multi-model spatial processor. Receives all detections (person + sink + "
    "hand) and active person tracks. Performs: hand->person assignment "
    "(point_inside_expanded_box + nearest center), person->sink assignment "
    "(normalised diagonal distance), hand-near-sink check (margin-expanded overlap), "
    "EMA movement tracking (deque=12, alpha=0.2), positive_frames / "
    "no_detection_frames counters, grace-period-aware inside_sink_zone. "
    "State resets on worker restart. evict_stale(3.0s) by timestamp.",
)

# ── Section 6.6: ComplianceRuleEngine ────────────────────────────────────────
i = find("_MAX_LOST_FRAMES=5: persons not seen for more than 5 frames receive a synthetic inside_sink_zone=False observation")
replace_para(
    i,
    "_MAX_LOST_FRAMES=5: persons in non-resting states not seen for >5 frames get a "
    "synthetic inside_sink_zone=False observation. source_fps threaded through to "
    "video_time_str() and evaluate() for MM:SS action logging.",
)

# ── Section 7.2: Observation data model ──────────────────────────────────────
i = find("@dataclass(frozen=True)\nclass Observation:\n    camera_id:       str\n    zone_id:         str\n    person_id:       int   ")
replace_para(
    i,
    "@dataclass(frozen=True)\n"
    "class Observation:\n"
    "    camera_id:            str\n"
    "    zone_id:              str\n"
    "    person_id:            int\n"
    "    timestamp:            datetime\n"
    "    frame_id:             int\n"
    "    inside_sink_zone:     bool     # hand near sink (with grace period)\n"
    "    hands_interacting:    bool     # confirmed: positive_frames >= threshold\n"
    "    hand_near_sink:       bool     # raw per-frame flag\n"
    "    hand_count:           int      # hands assigned to this person\n"
    "    hand_movement_score:  float    # EMA pixels/frame\n"
    "    positive_frames:      int      # consecutive frames with washing evidence\n"
    "    no_detection_frames:  int      # consecutive frames without\n"
    "    track_bbox:           BoundingBox | None\n"
    "    is_synthetic:         bool     # True for absence-synthesised observations",
)

# ── Section 7.3: TemporalMetrics ─────────────────────────────────────────────
i = find("@dataclass(frozen=True)\nclass TemporalMetrics:\n    at_sink_duration:       float   # Total seconds in AT_SINK state")
replace_para(
    i,
    "@dataclass(frozen=True)\n"
    "class TemporalMetrics:\n"
    "    at_sink_duration:      float   # seconds in NEAR_SINK_NOT_WASHING\n"
    "    washing_duration:      float   # seconds in WASHING (incl. current)\n"
    "    positive_frames:       int     # consecutive frames with washing evidence\n"
    "    no_detection_frames:   int     # consecutive frames without\n"
    "    hand_movement_score:   float   # EMA hand movement (pixels/frame)",
)

# ── Section 19.1: ByteTrack config example ────────────────────────────────────
i = find('models:\n  handwashing:\n    use_case: handwashing\n    path: "artifacts/models/handwash_detection.pt"')
replace_para(
    i,
    'models:\n'
    '  handwashing:\n'
    '    use_case: handwashing\n'
    '    device: cpu\n'
    '    image_size: 640\n'
    '    person:\n'
    '      path: "artifacts/models/yolo11n.pt"\n'
    '      confidence: 0.25\n'
    '      tracker: "bytetrack.yaml"   # ByteTrack on person model only\n'
    '    sink:\n'
    '      path: "artifacts/models/sink_detection.pt"\n'
    '      confidence: 0.50\n'
    '    hand:\n'
    '      path: "artifacts/models/handwash_detection.pt"\n'
    '      confidence: 0.50\n'
    '      presence_confidence: 0.25',
)

# ── Section 8.1: config.yaml snapshot ─────────────────────────────────────────
i = find("# BH_CV Compliance Detector — runtime configuration\n\nstream:\n  type: video                        # ")
replace_para(
    i,
    "# BH_CV Compliance Detector — runtime configuration\n\n"
    "stream:\n"
    "  type: video\n"
    "  camera_id: handwash-camera-01\n"
    "  sampling:\n"
    "    no_person_fps: 5.0\n"
    "    person_fps: 10.0\n\n"
    "models:\n"
    "  handwashing:\n"
    "    use_case: handwashing\n"
    "    device: cpu\n"
    "    image_size: 640\n"
    "    person:\n"
    "      path: artifacts/models/yolo11n.pt\n"
    "      confidence: 0.25\n"
    "      tracker: bytetrack.yaml\n"
    "    sink:\n"
    "      path: artifacts/models/sink_detection.pt\n"
    "      confidence: 0.50\n"
    "    hand:\n"
    "      path: artifacts/models/handwash_detection.pt\n"
    "      confidence: 0.50\n"
    "      presence_confidence: 0.25\n\n"
    "zones:\n"
    "  handwash_zone:\n"
    "    models: [handwashing.person, handwashing.sink, handwashing.hand]\n\n"
    "compliance:\n"
    "  trigger_class: hand_washing\n"
    "  handwash:\n"
    "    minimum_washing_duration_seconds: 5.0\n"
    "    absence_threshold_seconds: 5.0\n"
    "    evidence_buffer_max: 10\n"
    "  association:\n"
    "    person_box_expand_x: 0.15\n"
    "    person_box_expand_y: 0.20\n"
    "    max_person_sink_distance: 2.0\n"
    "    sink_horizontal_margin: 0.30\n"
    "    sink_top_margin: 1.50\n"
    "    sink_bottom_margin: 0.20\n"
    "  washing:\n"
    "    min_hands: 1\n"
    "    confirmation_frames: 3\n"
    "    reset_after_no_detection_seconds: 1.0\n"
    "    movement_history_size: 12\n"
    "    min_hand_movement: 2.0\n"
    "    require_hand_movement: true\n\n"
    "debug:\n"
    "  video_saving:\n"
    "    enabled: true\n"
    "    fallback_fps: 5.0\n"
    "    filename: detected.mp4",
)

# ── Update states table (TABLE 2) ─────────────────────────────────────────────
t = doc.tables[2]
new_states = [
    ("State", "Meaning"),
    ("NOT_NEAR_SINK", "Person not near any sink, or newly detected. Initial and resting state."),
    ("NEAR_SINK_NOT_WASHING", "Person hand is at the sink; washing not yet confirmed (positive_frames < threshold)."),
    ("WASHING", "Confirmed active washing: positive_frames >= confirmation_frames (3)."),
    ("WASHED", "Terminal: person left after sufficient washing duration (>= 5.0s). Set by RuleEvaluator."),
    ("NOT_WASHED", "Terminal: person left without sufficient washing. Violation. Set by RuleEvaluator."),
]
# Resize table if needed
while len(t.rows) > len(new_states):
    tr = t.rows[-1]._tr
    tr.getparent().remove(tr)
while len(t.rows) < len(new_states):
    t.add_row()
for r_idx, (state, meaning) in enumerate(new_states):
    t.rows[r_idx].cells[0].text = state
    t.rows[r_idx].cells[1].text = meaning

# ── Update transition table (TABLE 3) ────────────────────────────────────────
t3 = doc.tables[3]
new_transitions = [
    ("Current State", "Observation Condition", "Next State", "Reason"),
    ("WASHED / NOT_WASHED", "any (new detection)", "NOT_NEAR_SINK", "new_visit_after_terminal"),
    ("ANY", "inside_sink_zone=False", "NOT_NEAR_SINK", "left_sink_area"),
    ("NOT_NEAR_SINK", "inside_sink_zone=False", "NOT_NEAR_SINK", "not_near_sink"),
    ("NOT_NEAR_SINK", "inside_sink_zone=True", "NEAR_SINK_NOT_WASHING", "entered_sink_area"),
    ("NEAR_SINK_NOT_WASHING", "hands_interacting=True", "WASHING", "washing_confirmed"),
    ("NEAR_SINK_NOT_WASHING", "hands_interacting=False", "NEAR_SINK_NOT_WASHING", "near_sink_not_washing"),
    ("WASHING", "hands_interacting=False", "NEAR_SINK_NOT_WASHING", "washing_stopped"),
    ("WASHING", "hands_interacting=True", "WASHING", "still_washing"),
    ("NOT_NEAR_SINK (exit)", "RuleEvaluator: washing_seconds >= 5.0s", "WASHED", "compliant_exit"),
    ("NOT_NEAR_SINK (exit)", "RuleEvaluator: washing_seconds < 5.0s", "NOT_WASHED", "violation"),
]
while len(t3.rows) > len(new_transitions):
    tr = t3.rows[-1]._tr
    tr.getparent().remove(tr)
while len(t3.rows) < len(new_transitions):
    t3.add_row()
for r_idx, row_data in enumerate(new_transitions):
    for c_idx, cell_text in enumerate(row_data):
        if c_idx < len(t3.rows[r_idx].cells):
            t3.rows[r_idx].cells[c_idx].text = cell_text

# ── Update ObservationBuilder params table (TABLE 6) ─────────────────────────
t6 = doc.tables[6]
new_obs_params = [
    ("Parameter", "Default", "Description"),
    ("person_class", "person", "YOLO class name for people"),
    ("sink_class", "sink", "YOLO class name for sinks"),
    ("hand_class", "hand_washing", "YOLO class name for hand/washing activity (trigger_class)"),
    ("person_box_expand_x", "0.15", "Horizontal expansion ratio for hand->person box assignment"),
    ("person_box_expand_y", "0.20", "Vertical expansion ratio for hand->person box assignment"),
    ("max_person_sink_distance", "2.0", "Max normalised distance (person lower-center to sink center)"),
    ("sink_horizontal_margin", "0.30", "Lateral margin around sink for hand proximity check"),
    ("sink_top_margin", "1.50", "Upward margin (hands reach above the sink bowl)"),
    ("sink_bottom_margin", "0.20", "Downward margin for hand proximity check"),
    ("min_hands_for_washing", "1", "Minimum assigned hands to count as washing evidence"),
    ("min_hand_movement", "2.0", "Minimum EMA movement score (pixels/frame)"),
    ("require_hand_movement", "true", "If True, movement required for washing evidence"),
    ("confirmation_frames", "3", "Consecutive positive frames before hands_interacting=True"),
    ("reset_after_no_detection_seconds", "1.0", "Grace period before treating person as having left sink"),
    ("movement_history_size", "12", "Rolling window size for hand center history (deque maxlen)"),
]
while len(t6.rows) > len(new_obs_params):
    tr = t6.rows[-1]._tr
    tr.getparent().remove(tr)
while len(t6.rows) < len(new_obs_params):
    t6.add_row()
for r_idx, row_data in enumerate(new_obs_params):
    for c_idx, cell_text in enumerate(row_data):
        if c_idx < len(t6.rows[r_idx].cells):
            t6.rows[r_idx].cells[c_idx].text = cell_text

# ── Update config parameters table (TABLE 9) ─────────────────────────────────
t9 = doc.tables[9]
new_cfg_params = [
    ("Parameter", "Type", "Default", "Description"),
    ("stream.type", "string", "video", "Source type: 'video' or 'rtsp'"),
    ("stream.camera_id", "string", "-", "Must match cameras section key"),
    ("stream.sampling.no_person_fps", "float", "5.0", "Processing FPS when no person detected"),
    ("stream.sampling.person_fps", "float", "10.0", "Processing FPS when person is present"),
    ("stream.buffer.max_size", "int", "1", "Max frames buffered before eviction"),
    ("stream.reconnect.initial_delay_seconds", "float", "1.0", "Initial RTSP reconnect delay"),
    ("stream.reconnect.max_delay_seconds", "float", "15.0", "Maximum RTSP reconnect delay"),
    ("models.handwashing.person.confidence", "float", "0.25", "Person detection confidence threshold"),
    ("models.handwashing.sink.confidence", "float", "0.50", "Sink detection confidence threshold"),
    ("models.handwashing.hand.confidence", "float", "0.50", "Hand/washing confidence threshold (active washing)"),
    ("models.handwashing.hand.presence_confidence", "float", "0.25", "Hand presence threshold (detect near sink)"),
    ("models.<use_case>.device", "string", "cpu", "Inference device (cpu or cuda)"),
    ("models.<use_case>.image_size", "int", "640", "YOLO input resolution"),
    ("inference.queue_size", "int", "10", "Max frames queued for inference"),
    ("inference.warmup", "bool", "true", "Run warmup forward pass at startup"),
    ("inference.thumbnail.max_width", "int", "320", "Thumbnail resize width (pixels)"),
    ("inference.thumbnail.jpeg_quality", "int", "70", "JPEG compression quality [1,100]"),
    ("compliance.trigger_class", "string", "hand_washing", "YOLO class for active handwash"),
    ("compliance.handwash.minimum_washing_duration_seconds", "float", "5.0", "Required cumulative WASHING time for compliance"),
    ("compliance.handwash.absence_threshold_seconds", "float", "5.0", "Retained for compat; not used in exit-based evaluator"),
    ("compliance.handwash.evidence_buffer_max", "int", "10", "Max thumbnails per violation episode"),
    ("compliance.association.person_box_expand_x", "float", "0.15", "Hand-to-person box horizontal expansion"),
    ("compliance.association.person_box_expand_y", "float", "0.20", "Hand-to-person box vertical expansion"),
    ("compliance.association.max_person_sink_distance", "float", "2.0", "Max normalised person-to-sink distance"),
    ("compliance.association.sink_horizontal_margin", "float", "0.30", "Hand proximity lateral margin (x sink width)"),
    ("compliance.association.sink_top_margin", "float", "1.50", "Hand proximity upward margin (x sink height)"),
    ("compliance.association.sink_bottom_margin", "float", "0.20", "Hand proximity downward margin"),
    ("compliance.washing.min_hands", "int", "1", "Minimum hands for washing evidence"),
    ("compliance.washing.confirmation_frames", "int", "3", "Frames before hands_interacting=True"),
    ("compliance.washing.reset_after_no_detection_seconds", "float", "1.0", "Grace period before zone exit"),
    ("compliance.washing.movement_history_size", "int", "12", "Rolling window for hand movement"),
    ("compliance.washing.min_hand_movement", "float", "2.0", "Min EMA movement (pixels/frame)"),
    ("compliance.washing.require_hand_movement", "bool", "true", "Movement required for washing evidence"),
    ("debug.video_saving.enabled", "bool", "true", "Record annotated MP4 video clip"),
    ("debug.video_saving.fallback_fps", "float", "5.0", "FPS when source FPS is unknown"),
    ("debug.video_saving.filename", "string", "detected.mp4", "Output filename (mp4v codec)"),
]
while len(t9.rows) > len(new_cfg_params):
    tr = t9.rows[-1]._tr
    tr.getparent().remove(tr)
while len(t9.rows) < len(new_cfg_params):
    t9.add_row()
for r_idx, row_data in enumerate(new_cfg_params):
    for c_idx, cell_text in enumerate(row_data):
        if c_idx < len(t9.rows[r_idx].cells):
            t9.rows[r_idx].cells[c_idx].text = cell_text

# ── Update HandwashRuleConfig params table (TABLE 8) ─────────────────────────
t8 = doc.tables[8]
new_rule_params = [
    ("Config Field", "Type", "Description"),
    ("minimum_washing_duration_seconds", "float", "Required cumulative WASHING time for compliant exit (5.0s)"),
    ("absence_threshold_seconds", "float", "Retained for compat; not used in exit-based logic"),
    ("evidence_buffer_max", "int", "Max thumbnails buffered per violation episode (ring buffer)"),
    ("confirmation_frames", "int", "Consecutive positive frames before hands_interacting=True"),
    ("reset_after_no_detection_seconds", "float", "Seconds without hand-near-sink before zone exit"),
    ("min_hands_for_washing", "int", "Minimum associated hands required for washing evidence"),
    ("min_hand_movement", "float", "Minimum EMA movement score (pixels/frame)"),
    ("require_hand_movement", "bool", "Movement is a required condition for washing evidence"),
    ("person_box_expand_x / _y", "float", "Person bbox expansion for hand assignment"),
    ("max_person_sink_distance", "float", "Max normalised person-to-sink distance"),
    ("sink_horizontal_margin", "float", "Lateral margin around sink for hand proximity"),
    ("sink_top_margin", "float", "Upward margin (1.50 x sink height - hands reach above)"),
    ("sink_bottom_margin", "float", "Downward margin for hand proximity check"),
]
while len(t8.rows) > len(new_rule_params):
    tr = t8.rows[-1]._tr
    tr.getparent().remove(tr)
while len(t8.rows) < len(new_rule_params):
    t8.add_row()
for r_idx, row_data in enumerate(new_rule_params):
    for c_idx, cell_text in enumerate(row_data):
        if c_idx < len(t8.rows[r_idx].cells):
            t8.rows[r_idx].cells[c_idx].text = cell_text

# ── Save ──────────────────────────────────────────────────────────────────────
doc.save(PATH)
print("DONE — document saved to", PATH)
